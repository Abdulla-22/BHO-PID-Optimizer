import math
import os
import time
import threading
import traceback
import queue
import ctypes
from dataclasses import dataclass
from datetime import datetime
from typing import Callable, Dict, List, Optional, Tuple

import numpy as np
import pandas as pd
# Make Windows report the real screen size before Tkinter creates the GUI.
# This prevents fullscreen offset and wrong scaling on 1920x1080 displays.
try:
    ctypes.windll.shcore.SetProcessDpiAwareness(2)  # Per-monitor DPI aware
except Exception:
    try:
        ctypes.windll.user32.SetProcessDPIAware()
    except Exception:
        pass

import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import customtkinter as ctk

from scipy import signal

from PIL import Image, ImageTk

import matplotlib
matplotlib.use("TkAgg")
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
from matplotlib.figure import Figure

# Set CustomTkinter theme
ctk.set_appearance_mode("dark")
ctk.set_default_color_theme("blue")

# ============================================================
# Data structures
# ============================================================

@dataclass
class Plant:
    name: str
    num: np.ndarray
    den: np.ndarray
    step_amp: float

@dataclass
class StepInfo:
    overshoot: float
    rise_time: float
    settling_time: float
    steady_state_error_percent: float

# ============================================================
# Utility functions & Math Logic
# ============================================================

BIG_COST = 1e8

def as_array(x) -> np.ndarray:
    return np.asarray(x, dtype=float).flatten()

def pad_left(a: np.ndarray, n: int) -> np.ndarray:
    a = as_array(a)
    if len(a) >= n: return a
    return np.concatenate([np.zeros(n - len(a)), a])

def safe_float(value, default=0.0) -> float:
    try: return float(value)
    except Exception: return default

def parse_coefficients(text: str) -> np.ndarray:
    cleaned = text.replace(",", " ").replace(";", " ").strip()
    if not cleaned: raise ValueError("Coefficient field is empty.")
    return np.asarray([float(v) for v in cleaned.split()], dtype=float)

def init_ballandbeam(params: Dict[str, float]) -> Plant:
    m, R, g, L, d, J = params.get("m", 0.111), params.get("R", 0.015), params.get("g", 9.8), params.get("L", 1.0), params.get("d", 0.03), params.get("J", 9.99e-6)
    K = (m * g * d) / (L * ((J / (R ** 2)) + m))
    return Plant(name="BallAndBeam", num=np.array([K], dtype=float), den=np.array([1.0, 0.0, 0.0], dtype=float), step_amp=params.get("step_amp", 10.0))

def init_cruise(params: Dict[str, float]) -> Plant:
    m, b = params.get("m", 1000.0), params.get("b", 50.0)
    return Plant(name="CruiseControl", num=np.array([1.0], dtype=float), den=np.array([m, b], dtype=float), step_amp=params.get("step_amp", 10.0))

def init_motor(params: Dict[str, float]) -> Plant:
    J, b, K, R, L = params.get("J", 0.01), params.get("b", 1e-6), params.get("K", 0.1), params.get("R", 2.18), params.get("L", 2.3e-3)
    return Plant(name="DCMotorSpeed", num=np.array([K], dtype=float), den=np.array([(J * L), ((J * R) + (L * b)), ((b * R) + (K ** 2))], dtype=float), step_amp=params.get("step_amp", 50.0))

def init_custom_tf(params: Dict[str, float]) -> Plant:
    num, den = as_array(params["num"]), as_array(params["den"])
    if len(num) == 0 or np.allclose(num, 0) or len(den) == 0 or np.allclose(den, 0): raise ValueError("Custom coefficients cannot be empty or zero.")
    return Plant(name="CustomTF", num=num, den=den, step_amp=params.get("step_amp", 1.0))

def expand_controller_gains(k_opt: np.ndarray, controller_type: str) -> np.ndarray:
    k_opt = as_array(k_opt)
    ctype = controller_type.upper()
    if ctype == "PI": return np.array([k_opt[0], k_opt[1], 0.0], dtype=float)
    if ctype == "PD": return np.array([k_opt[0], 0.0, k_opt[1]], dtype=float)
    if ctype == "PID": return np.array([k_opt[0], k_opt[1], k_opt[2]], dtype=float)
    raise ValueError("Invalid controller type.")

def controller_tf(k_full: np.ndarray, controller_type: str) -> Tuple[np.ndarray, np.ndarray]:
    kp, ki, kd = as_array(k_full)
    ctype = controller_type.upper()
    if ctype == "PI": return np.array([kp, ki], dtype=float), np.array([1.0, 0.0], dtype=float)
    if ctype == "PD": return np.array([kd, kp], dtype=float), np.array([1.0], dtype=float)
    if ctype == "PID": return np.array([kd, kp, ki], dtype=float), np.array([1.0, 0.0], dtype=float)
    raise ValueError("Invalid controller type.")

def closed_loop_tf(plant: Plant, k_full: np.ndarray, controller_type: str) -> signal.TransferFunction:
    c_num, c_den = controller_tf(k_full, controller_type)
    open_num, open_den = np.polymul(c_num, plant.num), np.polymul(c_den, plant.den)
    n = max(len(open_num), len(open_den))
    closed_den = np.trim_zeros(pad_left(open_den, n) + pad_left(open_num, n), "f")
    closed_num = np.trim_zeros(pad_left(open_num, n), "f")
    if len(closed_num) == 0 or len(closed_den) == 0: raise ValueError("Invalid closed-loop transfer function.")
    return signal.TransferFunction(closed_num, closed_den)

def simulate_feedback_response(plant: Plant, k_full: np.ndarray, controller_type: str, sim_time: float, points: int = 300) -> Tuple[np.ndarray, np.ndarray]:
    sys_cl = closed_loop_tf(plant, k_full, controller_type)
    t = np.linspace(0.0, sim_time, points)
    tout, y = signal.step(sys_cl, T=t)
    y = plant.step_amp * np.asarray(y, dtype=float)
    if not np.all(np.isfinite(y)): raise ValueError("Simulation produced non-finite values.")
    return np.asarray(tout, dtype=float), y

def compute_step_info(t: np.ndarray, y: np.ndarray, ref: float) -> StepInfo:
    t, y = as_array(t), as_array(y)
    eps = np.finfo(float).eps
    if len(t) == 0 or len(y) == 0 or len(t) != len(y) or not np.all(np.isfinite(y)): return StepInfo(np.nan, np.nan, np.nan, np.nan)
    ref_abs = abs(ref) if abs(ref) > eps else 1.0
    overshoot = max(0.0, ((np.max(y) - ref) / abs(ref)) * 100.0) if abs(ref) > eps else 0.0
    ess_percent = (abs(ref - y[-1]) / abs(ref)) * 100.0 if abs(ref) > eps else abs(y[-1])
    try:
        idx10 = np.where(y >= 0.10 * ref)[0] if ref >= 0 else np.where(y <= 0.10 * ref)[0]
        idx90 = np.where(y >= 0.90 * ref)[0] if ref >= 0 else np.where(y <= 0.90 * ref)[0]
        rise_time = float(t[idx90[0]] - t[idx10[0]]) if len(idx10) > 0 and len(idx90) > 0 and float(t[idx90[0]] - t[idx10[0]]) >= 0 else np.nan
    except: rise_time = np.nan
    outside = np.where(np.abs(y - ref) > 0.02 * ref_abs)[0]
    settling_time = 0.0 if len(outside) == 0 else (float(t[outside[-1] + 1]) if outside[-1] < len(t) - 1 else np.nan)
    return StepInfo(float(overshoot), float(rise_time), float(settling_time), float(ess_percent))

def generic_cost(k_full, plant, ctype, sim_time, opt_mode, w_os, w_rt, w_ess, weights, penalty_gain) -> float:
    time.sleep(0.001)  # Yield GIL for UI responsiveness
    try:
        t, y = simulate_feedback_response(plant, k_full, ctype, sim_time)
        if len(y) == 0 or len(t) == 0 or not np.isfinite(y[-1]) or not np.all(np.isfinite(y)): return BIG_COST
        info = compute_step_info(t, y, plant.step_amp)
        if not np.isfinite(info.overshoot) or not np.isfinite(info.rise_time): return 1e7
        ref = plant.step_amp
        iae_norm = np.trapezoid(np.abs(ref - y), t) / (abs(ref) * sim_time) if abs(ref) >= 1e-12 else np.trapezoid(np.abs(ref - y), t) / max(sim_time, 1e-12)
        actual_ess = (abs(ref - y[-1]) / abs(ref)) * 100.0 if abs(ref) >= 1e-12 else abs(y[-1]) * 100.0
        w1, w2, w3, w4 = weights
        if opt_mode == 1:
            cost = (w1 * iae_norm) + (w2 * info.overshoot / 100.0) + (w3 * info.rise_time / max(sim_time, 1e-12)) + (w4 * actual_ess / 100.0)
            return float(cost) if np.isfinite(cost) else BIG_COST
        e_os, e_rt, e_e = abs(info.overshoot - w_os)/max(w_os,1e-6), abs(info.rise_time - w_rt)/max(w_rt,1e-6), abs(actual_ess - w_ess)/max(w_ess,1e-6)
        cost = (w1 * iae_norm) + (w2 * e_os) + (w3 * e_rt) + (w4 * e_e)
        if info.overshoot > w_os: cost += penalty_gain * ((info.overshoot - w_os) / max(w_os, 1e-6))
        if info.rise_time > w_rt: cost += penalty_gain * ((info.rise_time - w_rt) / max(w_rt, 1e-6))
        if actual_ess > w_ess: cost += penalty_gain * ((actual_ess - w_ess) / max(w_ess, 1e-6))
        return float(cost) if np.isfinite(cost) else BIG_COST
    except Exception: return BIG_COST

def cost_ballandbeam(k, p, c, t, o, w1, w2, w3): return generic_cost(k, p, c, t, o, w1, w2, w3, (0.20, 0.30, 0.25, 0.25), 8.0)
def cost_cruise(k, p, c, t, o, w1, w2, w3): return generic_cost(k, p, c, t, o, w1, w2, w3, (0.15, 0.30, 0.25, 0.30), 5.0)
def cost_motor(k, p, c, t, o, w1, w2, w3): return generic_cost(k, p, c, t, o, w1, w2, w3, (0.10, 0.35, 0.35, 0.20), 5.0)
def cost_custom(k, p, c, t, o, w1, w2, w3): return generic_cost(k, p, c, t, o, w1, w2, w3, (0.20, 0.30, 0.30, 0.20), 5.0)

def black_hole_algorithm(cost_fcn, n_pop, max_iter, var_min, var_max, controller_type, plant, sim_time, stop_fcn=None, progress_fcn=None):
    var_min, var_max = as_array(var_min), as_array(var_max)
    n_var = len(var_min)
    if stop_fcn is None: stop_fcn = lambda: False
    rng = np.random.default_rng()
    positions = var_min + rng.random((n_pop, n_var)) * (var_max - var_min)
    costs = np.array([cost_fcn(p) for p in positions])
    history_rows = []

    for it in range(1, max_iter + 1):
        if stop_fcn(): return None, pd.DataFrame(history_rows)
        idx_bh = int(np.argmin(costs))
        bh_pos = positions[idx_bh, :].copy()
        
        moved_positions, moved_costs = positions.copy(), costs.copy()
        for i in range(n_pop):
            if stop_fcn(): return None, pd.DataFrame(history_rows)
            if i != idx_bh:
                moved_positions[i, :] += rng.random(n_var) * (bh_pos - moved_positions[i, :])
                moved_positions[i, :] = np.clip(moved_positions[i, :], var_min, var_max)
                moved_costs[i] = cost_fcn(moved_positions[i, :])

        idx_bh_new = int(np.argmin(moved_costs))
        bh_pos, bh_cost = moved_positions[idx_bh_new, :].copy(), moved_costs[idx_bh_new]
        sum_costs = np.sum(moved_costs)
        radius = 0.0 if sum_costs == 0 else bh_cost / sum_costs

        for i in range(n_pop):
            if stop_fcn(): return None, pd.DataFrame(history_rows)
            if i != idx_bh_new and np.linalg.norm(bh_pos - moved_positions[i, :]) < radius:
                moved_positions[i, :] = var_min + rng.random(n_var) * (var_max - var_min)
                moved_costs[i] = cost_fcn(moved_positions[i, :])

        positions, costs = moved_positions, moved_costs
        idx_best = int(np.argmin(costs))
        k_full = expand_controller_gains(positions[idx_best, :], controller_type)
        
        ov, rt, ess = np.nan, np.nan, np.nan
        try:
            t_r, y_r = simulate_feedback_response(plant, k_full, controller_type, sim_time)
            info = compute_step_info(t_r, y_r, plant.step_amp)
            ov, rt, ess = info.overshoot, info.rise_time, info.steady_state_error_percent
        except: pass

        history_rows.append({"Iteration": it, "Kp": k_full[0], "Ki": k_full[1], "Kd": k_full[2], "Cost": costs[idx_best], "Overshoot": ov, "Rise Time": rt, "Steady-State Error": ess})
        
        if progress_fcn is not None: progress_fcn(history_rows[-1])
        time.sleep(0.1)  # Smooth delay so user can see iterations one by one

    return positions[int(np.argmin(costs)), :].copy(), pd.DataFrame(history_rows)


# ============================================================
# GUI Application (CustomTkinter)
# ============================================================

class ImageStagePlayer(ctk.CTkFrame):
    def __init__(self, parent, bg_color="#ffffff", **kwargs):
        super().__init__(parent, fg_color=bg_color, **kwargs)
        self.frames = []
        self.index = 0
        self.job = None
        self._load_id = 0
        self._cache = {}
        self.label = tk.Label(self, bg=bg_color, fg="#1f538d", text="Black Hole Algorithm Visualization", font=("Segoe UI", 16, "bold"))
        self.label.pack(fill="both", expand=True)

    def preload(self, paths_list, w=220, h=220):
        threading.Thread(target=self._preload_worker, args=(paths_list, w, h), daemon=True).start()

    def _preload_worker(self, paths_list, w, h):
        for paths in paths_list:
            cache_key = tuple(paths)
            if cache_key in self._cache: continue
            pil_frames = []
            for path in paths:
                if path and os.path.isfile(path):
                    ext = os.path.splitext(path)[1].lower()
                    try:
                        img = Image.open(path)
                        if ext == ".gif":
                            i = 0
                            while True:
                                try:
                                    img.seek(i)
                                    frame = img.copy().convert("RGBA")
                                    frame.thumbnail((w, h), Image.Resampling.BILINEAR)
                                    pil_frames.append(frame)
                                    i += 1
                                except EOFError: break
                        else:
                            img = img.convert("RGBA")
                            img.thumbnail((w, h), Image.Resampling.BILINEAR)
                            pil_frames.append(img)
                    except Exception: pass
            self._cache[cache_key] = pil_frames

    def show(self, paths, animate=False):
        self.stop()
        self.frames = []
        if isinstance(paths, str):
            paths = [paths]
            
        self.update_idletasks()
        w = self.winfo_width()
        h = self.winfo_height()
        if w <= 1: w = 220
        if h <= 1: h = 220

        cache_key = tuple(paths)
        if cache_key in self._cache:
            self._apply_frames(self._cache[cache_key], animate, None)
        else:
            self._load_id += 1
            curr_id = self._load_id
            threading.Thread(target=self._load_async, args=(paths, w, h, animate, curr_id), daemon=True).start()

    def _load_async(self, paths, w, h, animate, load_id):
        pil_frames = []
        for path in paths:
            if path and os.path.isfile(path):
                ext = os.path.splitext(path)[1].lower()
                try:
                    img = Image.open(path)
                    if ext == ".gif":
                        i = 0
                        while True:
                            try:
                                img.seek(i)
                                frame = img.copy().convert("RGBA")
                                frame.thumbnail((w, h), Image.Resampling.BILINEAR)
                                pil_frames.append(frame)
                                i += 1
                                if i % 10 == 0: time.sleep(0.001)
                            except EOFError: break
                    else:
                        img = img.convert("RGBA")
                        img.thumbnail((w, h), Image.Resampling.BILINEAR)
                        pil_frames.append(img)
                except Exception: pass
        self._cache[tuple(paths)] = pil_frames
        self.after(0, lambda: self._apply_frames(pil_frames, animate, load_id))

    def _apply_frames(self, pil_frames, animate, load_id):
        if load_id is not None and load_id != self._load_id: return
        self.frames = [ImageTk.PhotoImage(img) for img in pil_frames]
        if self.frames:
            self.label.configure(image=self.frames[0], text="")
            if animate and len(self.frames) > 1: self.play(80)
        else:
            self.label.configure(image="", text="")

    def play(self, delay_ms=80):
        self.stop()
        if len(self.frames) > 1:
            self.delay_ms = delay_ms
            self._animate()

    def stop(self):
        if self.job:
            try: self.after_cancel(self.job)
            except: pass
            self.job = None

    def _animate(self):
        if not self.frames: return
        self.index = (self.index + 1) % len(self.frames)
        self.label.configure(image=self.frames[self.index])
        self.job = self.after(self.delay_ms, self._animate)

    def clear(self, text="Ready"):
        self.stop()
        self.label.configure(image="", text=text)


class BlackHoleOptimizerApp(ctk.CTk):
    def __init__(self):
        super().__init__()
        self.title("Black Hole Optimizer")
        
        window_w = 1450
        window_h = 720
        screen_w = self.winfo_screenwidth()
        screen_h = self.winfo_screenheight()
        
        if window_w > screen_w: window_w = screen_w - 50
        if window_h > screen_h: window_h = screen_h - 50
        
        x = int((screen_w - window_w) / 2)
        y = int((screen_h - window_h) / 2)
        
        self.geometry(f"{window_w}x{window_h}+-8+-2")
        self.minsize(1200, 720)
        
        # Deep space MATLAB-like colors
        self.col_bg = "#06111f"
        self.col_panel = "#0b2135"
        self.col_border = "#164c78"
        self.col_accent = "#38bdf8"
        self.configure(fg_color=self.col_bg)

        self.stop_requested = False
        self.is_running = False
        self.last_history = pd.DataFrame()
        self.last_iter_plotted = 0
        self.app_dir = os.path.dirname(os.path.abspath(__file__))
        self.images_dir = os.path.join(self.app_dir, "images")

        self.stage_images = {
            "generation": "1 stars generation.gif",
            "running": ["First.gif", "Second.gif"],
            "finish": "5.1.2 finish case 1.png"
        }

        self._setup_ttk_styles_for_treeview()
        self._build_ui()
        self._set_initial_state()
        
        # Preload all stages for instant image transitions
        all_paths = []
        for k, v in self.stage_images.items():
            paths = v if isinstance(v, list) else [v]
            paths = [os.path.join(self.images_dir, p) for p in paths]
            all_paths.append(paths)
        self.viz_player.preload(all_paths)
        
        # Fullscreen handling
        self.is_fullscreen = False
        self.normal_geometry = self.geometry()
        self._geometry_restore_job = None

        self.bind_all("<F11>", self.toggle_fullscreen)
        self.bind_all("<Escape>", self.exit_fullscreen)
        self.bind("<Configure>", self._remember_normal_geometry)

    def toggle_fullscreen(self, event=None):
        if self.is_fullscreen:
            self.exit_fullscreen()
        else:
            self.enter_fullscreen()
        return "break"

    def enter_fullscreen(self, event=None):
        """
        Enter true fullscreen.

        This uses native Tk fullscreen after DPI awareness is set before Tkinter import.
        No overrideredirect is used because it causes offset and layout issues.
        """
        if self.is_fullscreen:
            return "break"

        self.update_idletasks()
        self.normal_geometry = self.geometry()
        self.is_fullscreen = True

        # Reset any old window-manager states.
        self.overrideredirect(False)
        self.attributes("-topmost", False)
        self.state("normal")
        self.update_idletasks()

        # Native fullscreen should cover the full 1920x1080 display.
        self.attributes("-fullscreen", True)

        # Refresh layout after entering fullscreen.
        self.after(80, self._refresh_layout_after_fullscreen)
        return "break"

    def exit_fullscreen(self, event=None):
        """
        Exit fullscreen and restore the exact normal window size/position.
        """
        if not self.is_fullscreen:
            return "break"

        self.is_fullscreen = False

        self.attributes("-fullscreen", False)
        self.overrideredirect(False)
        self.attributes("-topmost", False)
        self.state("normal")

        if self._geometry_restore_job is not None:
            try:
                self.after_cancel(self._geometry_restore_job)
            except Exception:
                pass

        self._geometry_restore_job = self.after(80, self._restore_normal_geometry)
        return "break"

    def _restore_normal_geometry(self):
        try:
            if self.normal_geometry:
                self.geometry(self.normal_geometry)
            self.update_idletasks()
            self.lift()
            self.focus_force()
        except Exception:
            pass
        finally:
            self._geometry_restore_job = None

    def _refresh_layout_after_fullscreen(self):
        try:
            self.update_idletasks()

            # Force matplotlib canvases to redraw using the new fullscreen size.
            if hasattr(self, "canvas_tf"):
                self.canvas_tf.draw_idle()
            if hasattr(self, "canvas_cost"):
                self.canvas_cost.draw_idle()
            if hasattr(self, "canvas_sys"):
                self.canvas_sys.draw_idle()

            self.lift()
            self.focus_force()
        except Exception:
            pass

    def _remember_normal_geometry(self, event=None):
        """
        Save normal geometry only when the window is not fullscreen.
        This prevents fullscreen geometry from overwriting the restore size.
        """
        if self.is_fullscreen:
            return

        try:
            geometry = self.geometry()
            if geometry and "x" in geometry and "+" in geometry:
                self.normal_geometry = geometry
        except Exception:
            pass

    def _setup_ttk_styles_for_treeview(self):
        style = ttk.Style(self)
        style.theme_use("default")
        style.configure("Treeview", background=self.col_panel, foreground="white", fieldbackground=self.col_panel, borderwidth=0, rowheight=24)
        style.map('Treeview', background=[('selected', self.col_border)])
        style.configure("Treeview.Heading", background="#071827", foreground=self.col_accent, relief="flat", font=("Segoe UI", 9, "bold"))
        style.map("Treeview.Heading", background=[('active', self.col_border)])

    def _create_section(self, parent, title, row, col, rowspan=1, colspan=1, sticky="nsew"):
        frame = ctk.CTkFrame(parent, fg_color=self.col_panel, border_width=1, border_color=self.col_border, corner_radius=8)
        frame.grid(row=row, column=col, rowspan=rowspan, columnspan=colspan, sticky=sticky, padx=5, pady=5)
        lbl = ctk.CTkLabel(frame, text=title, text_color=self.col_accent, font=("Segoe UI", 12, "bold"))
        lbl.pack(anchor="w", padx=10, pady=(5, 0))
        content = ctk.CTkFrame(frame, fg_color="transparent")
        content.pack(fill="both", expand=True, padx=10, pady=(0, 5))
        return content

    def _build_ui(self):
        # Grid layout
        self.columnconfigure(0, weight=1)
        self.rowconfigure(1, weight=1)

        # Header
        header = ctk.CTkFrame(self, fg_color="transparent")
        header.grid(row=0, column=0, sticky="ew", padx=20, pady=(15, 0))
        ctk.CTkLabel(header, text="Black Hole Optimizer", font=("Segoe UI", 28, "bold"), text_color="white").pack(side="left")
        ctk.CTkLabel(header, text=" | Advanced PID Tuning using Black Hole Algorithm", font=("Segoe UI", 14), text_color="#8bb6d6").pack(side="left", padx=10, pady=(8,0))

        main_frame = ctk.CTkFrame(self, fg_color="transparent")
        main_frame.grid(row=1, column=0, sticky="nsew", padx=10, pady=10)
        main_frame.columnconfigure(0, weight=0, minsize=420)
        main_frame.columnconfigure(1, weight=1)
        main_frame.rowconfigure(0, weight=1)

        left_panel = ctk.CTkFrame(main_frame, fg_color="transparent")
        left_panel.grid(row=0, column=0, sticky="nsew", padx=(0, 10))
        left_panel.columnconfigure(0, weight=1)
        left_panel.rowconfigure(2, weight=1)

        # --- LEFT PANEL ---
        # 1. Selection
        sec1 = self._create_section(left_panel, "1  System Selection", 0, 0, sticky="new")
        sec1.columnconfigure(1, weight=1)
        
        ctk.CTkLabel(sec1, text="Select System:", font=("Segoe UI", 11)).grid(row=0, column=0, sticky="w", pady=2)
        self.system_var = ctk.StringVar(value="Cruise Control")
        sys_cb = ctk.CTkOptionMenu(sec1, variable=self.system_var, values=["Ball & Beam", "Cruise Control", "Motor Speed Control", "Custom TF"], command=self._show_params, height=24, font=("Segoe UI", 11))
        sys_cb.grid(row=0, column=1, sticky="ew", padx=(10, 0), pady=2)

        ctk.CTkLabel(sec1, text="Controller:", font=("Segoe UI", 11)).grid(row=1, column=0, sticky="w", pady=2)
        self.controller_var = ctk.StringVar(value="PID")
        ctrl_cb = ctk.CTkOptionMenu(sec1, variable=self.controller_var, values=["PI", "PD", "PID"], height=24, font=("Segoe UI", 11))
        ctrl_cb.grid(row=1, column=1, sticky="ew", padx=(10, 0), pady=2)

        # 2. Parameters
        sec2 = self._create_section(left_panel, "2  Parameters", 1, 0, sticky="new")
        self.param_panels, self.entries = {}, {}
        self._build_params(sec2)

        # 3. Settings
        sec3 = self._create_section(left_panel, "3  Optimization Settings", 2, 0, sticky="nsew")
        for i in range(2): sec3.columnconfigure(i, weight=1)
        self.npop_var, self.maxit_var, self.sim_time_var = ctk.StringVar(value="30"), ctk.StringVar(value="50"), ctk.StringVar(value="10")
        self._add_entry(sec3, "Population:", self.npop_var, 0, 0)
        self._add_entry(sec3, "Max Iter:", self.maxit_var, 0, 1)
        self._add_entry(sec3, "Sim Time:", self.sim_time_var, 1, 0)

        self.min_kp, self.max_kp = ctk.StringVar(value="0"), ctk.StringVar(value="1000")
        self.min_ki, self.max_ki = ctk.StringVar(value="0"), ctk.StringVar(value="500")
        self.min_kd, self.max_kd = ctk.StringVar(value="0"), ctk.StringVar(value="100")
        self._add_entry(sec3, "Min Kp:", self.min_kp, 2, 0); self._add_entry(sec3, "Max Kp:", self.max_kp, 2, 1)
        self._add_entry(sec3, "Min Ki:", self.min_ki, 3, 0); self._add_entry(sec3, "Max Ki:", self.max_ki, 3, 1)
        self._add_entry(sec3, "Min Kd:", self.min_kd, 4, 0); self._add_entry(sec3, "Max Kd:", self.max_kd, 4, 1)

        self.best_mode_var = ctk.BooleanVar(value=True)
        ctk.CTkCheckBox(sec3, text="Best Optimization Mode", variable=self.best_mode_var, command=self._update_mode, font=("Segoe UI", 11), checkbox_height=18, checkbox_width=18).grid(row=5, column=0, columnspan=2, sticky="w", pady=(5, 2))
        self.w_os, self.w_rt, self.w_ess = ctk.StringVar(value="10"), ctk.StringVar(value="1"), ctk.StringVar(value="1")
        self.e_os = self._add_entry(sec3, "Wanted OS (%):", self.w_os, 6, 0)
        self.e_rt = self._add_entry(sec3, "Wanted RT (s):", self.w_rt, 6, 1)
        self.e_ess = self._add_entry(sec3, "Wanted ESS (%):", self.w_ess, 7, 0)

        # Buttons
        act_frame = ctk.CTkFrame(left_panel, fg_color="transparent")
        act_frame.grid(row=3, column=0, sticky="ew", pady=(0, 5))
        act_frame.columnconfigure(0, weight=1); act_frame.columnconfigure(1, weight=1); act_frame.columnconfigure(2, weight=1)
        self.btn_run = ctk.CTkButton(act_frame, text="▶ Run", fg_color="#166534", hover_color="#22c55e", font=("Segoe UI", 12, "bold"), command=self.run_opt, height=28)
        self.btn_stop = ctk.CTkButton(act_frame, text="⏸ Stop", fg_color="#7f1d1d", hover_color="#ef4444", font=("Segoe UI", 12, "bold"), state="disabled", command=self.stop_opt, height=28)
        self.btn_reset = ctk.CTkButton(act_frame, text="↻ Reset", fg_color="#1f2937", hover_color="#334155", font=("Segoe UI", 12), command=self.reset_results, height=28)
        self.btn_run.grid(row=0, column=0, sticky="ew", padx=(0, 5))
        self.btn_stop.grid(row=0, column=1, sticky="ew", padx=5)
        self.btn_reset.grid(row=0, column=2, sticky="ew", padx=(5, 0))

        # --- RIGHT PANEL ---
        right_frame = ctk.CTkFrame(main_frame, fg_color="transparent")
        right_frame.grid(row=0, column=1, sticky="nsew")
        right_frame.columnconfigure(0, weight=1)
        right_frame.rowconfigure(0, weight=0) # Viz area fixed
        right_frame.rowconfigure(1, weight=1) # Progress area stretches

        # Viz Area (RESIZED & REMOVED BUTTONS)
        viz_sec = self._create_section(right_frame, "Black Hole Algorithm Visualization & System TF", 0, 0)
        viz_sec.columnconfigure(0, weight=0); viz_sec.columnconfigure(1, weight=1)
        
        # Reduced height for ImageStagePlayer
        self.viz_player = ImageStagePlayer(viz_sec, bg_color="#ffffff", width=220, height=220)
        self.viz_player.grid(row=0, column=0, sticky="nsew", padx=(0, 10))
        self.viz_player.pack_propagate(False)
        self.viz_player.grid_propagate(False)

        right_viz_panel = ctk.CTkFrame(viz_sec, fg_color="transparent")
        right_viz_panel.grid(row=0, column=1, sticky="nsew")
        right_viz_panel.columnconfigure(0, weight=1)
        right_viz_panel.rowconfigure(1, weight=1)

        desc_box = ctk.CTkFrame(right_viz_panel, fg_color="#06111f", corner_radius=5)
        desc_box.grid(row=0, column=0, sticky="nsew", pady=(0, 5))
        ctk.CTkLabel(desc_box, text="Step Description", font=("Segoe UI", 11, "bold"), text_color=self.col_accent).pack(anchor="w", padx=10, pady=(5, 0))
        self.lbl_desc = ctk.CTkLabel(desc_box, text="Stars generated.", font=("Segoe UI", 11), text_color="white", wraplength=200, justify="left")
        self.lbl_desc.pack(anchor="w", padx=10, pady=2)
        ctk.CTkLabel(desc_box, text="Legend: ● Black Hole  ✦ Stars", font=("Segoe UI", 11), justify="left").pack(anchor="w", padx=10, pady=(5, 5))

        self.tf_box = ctk.CTkFrame(right_viz_panel, fg_color="#06111f", corner_radius=5)
        self.tf_box.grid(row=1, column=0, sticky="nsew")
        ctk.CTkLabel(self.tf_box, text="Transfer Function Preview", font=("Segoe UI", 11, "bold"), text_color=self.col_accent).pack(anchor="w", padx=10, pady=2)
        
        self.fig_tf = Figure(figsize=(3, 1), dpi=100, facecolor="#06111f")
        self.ax_tf = self.fig_tf.add_subplot(111)
        self.ax_tf.axis("off")
        self.canvas_tf = FigureCanvasTkAgg(self.fig_tf, master=self.tf_box)
        self.canvas_tf.get_tk_widget().pack(fill="both", expand=True, padx=5, pady=2)

        # Progress Area
        prog_sec = self._create_section(right_frame, "Optimization Progress", 1, 0)
        prog_sec.columnconfigure(0, weight=1); prog_sec.columnconfigure(1, weight=1); prog_sec.columnconfigure(2, weight=0)
        prog_sec.rowconfigure(0, weight=0); prog_sec.rowconfigure(1, weight=1)

        self.fig_cost = Figure(figsize=(4, 2.5), dpi=100, facecolor=self.col_panel)
        self.ax_cost = self.fig_cost.add_subplot(111)
        self.canvas_cost = FigureCanvasTkAgg(self.fig_cost, master=prog_sec)
        self.canvas_cost.get_tk_widget().grid(row=0, column=0, sticky="nsew", padx=(0, 5), pady=(0, 10))

        self.fig_sys = Figure(figsize=(4, 2.5), dpi=100, facecolor=self.col_panel)
        self.ax_sys = self.fig_sys.add_subplot(111)
        self.canvas_sys = FigureCanvasTkAgg(self.fig_sys, master=prog_sec)
        self.canvas_sys.get_tk_widget().grid(row=0, column=1, sticky="nsew", padx=5, pady=(0, 10))

        best_box = ctk.CTkFrame(prog_sec, fg_color="#06111f", corner_radius=5)
        best_box.grid(row=0, column=2, sticky="nsew", padx=(5, 0), pady=(0, 10))
        ctk.CTkLabel(best_box, text="Best Solution", font=("Segoe UI", 12, "bold"), text_color="white").pack(pady=(10, 10))
        self.metrics = {}
        for m, color in [("Kp", self.col_accent), ("Ki", self.col_accent), ("Kd", self.col_accent), ("Cost", "#ffd400"), ("Overshoot", "#ff4bd8"), ("Rise Time", self.col_accent), ("ESS", "#59ff45")]:
            f = ctk.CTkFrame(best_box, fg_color="transparent")
            f.pack(fill="x", padx=15, pady=1)
            ctk.CTkLabel(f, text=m, font=("Segoe UI", 11)).pack(side="left")
            lbl = ctk.CTkLabel(f, text="--", font=("Segoe UI", 11, "bold"), text_color=color)
            lbl.pack(side="right")
            self.metrics[m] = lbl

        tv_frame = ctk.CTkFrame(prog_sec, fg_color="transparent")
        tv_frame.grid(row=1, column=0, columnspan=3, sticky="nsew")
        tv_frame.columnconfigure(0, weight=1); tv_frame.rowconfigure(0, weight=1)
        cols = ["Iteration", "Kp", "Ki", "Kd", "Cost", "Overshoot", "Rise Time", "Steady-State Error"]
        self.tree = ttk.Treeview(tv_frame, columns=cols, show="headings")
        for c in cols: self.tree.heading(c, text=c); self.tree.column(c, width=100, anchor="center")
        self.tree.grid(row=0, column=0, sticky="nsew")
        sb = ttk.Scrollbar(tv_frame, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=sb.set); sb.grid(row=0, column=1, sticky="ns")

        # Footer
        footer = ctk.CTkFrame(self, fg_color=self.col_panel, height=50)
        footer.grid(row=2, column=0, sticky="ew")
        footer.columnconfigure(2, weight=1)
        ctk.CTkLabel(footer, text="Status: ", font=("Segoe UI", 12, "bold")).grid(row=0, column=0, padx=(20, 5), pady=10)
        self.lbl_status = ctk.CTkLabel(footer, text="Ready", font=("Segoe UI", 12, "bold"), text_color="#59ff45")
        self.lbl_status.grid(row=0, column=1, pady=10)
        self.lbl_prog = ctk.CTkLabel(footer, text="No run yet", text_color="#8bb6d6")
        self.lbl_prog.grid(row=0, column=2, sticky="w", padx=20)
        self.btn_save = ctk.CTkButton(footer, text="💾 Save Results", fg_color="#075985", hover_color="#0284c7", state="disabled", command=self.save_results)
        self.btn_save.grid(row=0, column=3, padx=20, pady=10)

    def _add_entry(self, parent, label, var, r, c):
        f = ctk.CTkFrame(parent, fg_color="transparent")
        f.grid(row=r, column=c, sticky="ew", padx=2, pady=2); f.columnconfigure(1, weight=1)
        ctk.CTkLabel(f, text=label, font=("Segoe UI", 11)).grid(row=0, column=0, sticky="w", padx=(0, 5))
        e = ctk.CTkEntry(f, textvariable=var, width=70, height=24, font=("Segoe UI", 11))
        e.grid(row=0, column=1, sticky="e")
        return e

    def _build_params(self, parent):
        parent.rowconfigure(0, weight=1)
        parent.columnconfigure(0, weight=1)
        for sys in ["Ball & Beam", "Cruise Control", "Motor Speed Control", "Custom TF"]:
            f = ctk.CTkFrame(parent, fg_color="transparent")
            f.grid(row=0, column=0, sticky="nsew")
            self.param_panels[sys] = f

        def v(name, val): self.entries[name] = ctk.StringVar(value=val); return self.entries[name]

        f1 = self.param_panels["Ball & Beam"]
        for i in range(2): f1.columnconfigure(i, weight=1)
        self._add_entry(f1, "m:", v("bb_m", "0.111"), 0, 0); self._add_entry(f1, "R:", v("bb_R", "0.015"), 0, 1)
        self._add_entry(f1, "g:", v("bb_g", "9.8"), 1, 0); self._add_entry(f1, "L:", v("bb_L", "1.0"), 1, 1)
        self._add_entry(f1, "d:", v("bb_d", "0.03"), 2, 0); self._add_entry(f1, "J:", v("bb_J", "9.99e-6"), 2, 1)
        self._add_entry(f1, "Step:", v("bb_s", "10"), 3, 0)

        f2 = self.param_panels["Cruise Control"]
        for i in range(2): f2.columnconfigure(i, weight=1)
        self._add_entry(f2, "m:", v("cr_m", "1000"), 0, 0); self._add_entry(f2, "b:", v("cr_b", "50"), 0, 1)
        self._add_entry(f2, "Step:", v("cr_s", "10"), 1, 0)

        f3 = self.param_panels["Motor Speed Control"]
        for i in range(2): f3.columnconfigure(i, weight=1)
        self._add_entry(f3, "J:", v("mo_J", "0.01"), 0, 0); self._add_entry(f3, "b:", v("mo_b", "1e-6"), 0, 1)
        self._add_entry(f3, "K:", v("mo_K", "0.1"), 1, 0); self._add_entry(f3, "R:", v("mo_R", "2.18"), 1, 1)
        self._add_entry(f3, "L:", v("mo_L", "2.3e-3"), 2, 0); self._add_entry(f3, "RPM:", v("mo_s", "50"), 2, 1)

        f4 = self.param_panels["Custom TF"]
        f4.columnconfigure(1, weight=1)
        self._add_entry(f4, "Num Coeffs:", v("ctf_n", "1"), 0, 0); self._add_entry(f4, "Den Coeffs:", v("ctf_d", "1 5"), 1, 0)
        self._add_entry(f4, "Step:", v("ctf_s", "1"), 2, 0)
        
        for var in self.entries.values():
            var.trace_add("write", lambda *args: self.after(50, self._update_tf_preview))

    def _show_params(self, choice=None):
        sel = self.system_var.get()
        for k, f in self.param_panels.items(): f.lift() if k == sel else f.lower()
        self.after(50, self._update_tf_preview)

    def _update_tf_preview(self, *args):
        if not hasattr(self, 'ax_tf'): return
        try:
            plant, _ = self._get_sys_data()
            def poly_latex(coeffs):
                if not len(coeffs): return "0"
                terms = []
                deg = len(coeffs) - 1
                for c in coeffs:
                    if c == 0 and deg > 0:
                        deg -= 1
                        continue
                    sgn = " - " if c < 0 and terms else ("-" if c < 0 else (" + " if terms else ""))
                    val = abs(c)
                    val_str = "" if (val == 1 and deg > 0) else (f"{int(val)}" if val.is_integer() else f"{val:g}")
                    var_str = "" if deg == 0 else ("s" if deg == 1 else f"s^{{{deg}}}")
                    terms.append(sgn + val_str + var_str)
                    deg -= 1
                return "".join(terms) if terms else "0"

            num_latex = poly_latex(plant.num)
            den_latex = poly_latex(plant.den)
            latex_str = f"$G(s) = \\frac{{{num_latex}}}{{{den_latex}}}$"
        except Exception:
            latex_str = "Invalid Parameters"

        self.ax_tf.clear()
        self.ax_tf.axis("off")
        self.ax_tf.text(0.5, 0.5, latex_str, fontsize=16, color="white", ha="center", va="center")
        self.canvas_tf.draw_idle()

    def _update_mode(self):
        st = "disabled" if self.best_mode_var.get() else "normal"
        self.e_os.configure(state=st); self.e_rt.configure(state=st); self.e_ess.configure(state=st)

    def _set_initial_state(self):
        self._show_params(); self._update_mode(); self._show_stage("blank")
        self._style_ax(self.ax_cost, "Cost vs Iteration", "Iteration", "Cost")
        self._style_ax(self.ax_sys, "System Response (Best Solution)", "Time (s)", "Amplitude")
        self.canvas_cost.draw(); self.canvas_sys.draw()

    def _style_ax(self, ax, title, xl, yl):
        ax.clear()
        ax.set_facecolor(self.col_panel)
        ax.set_title(title, color="white", fontsize=10, fontweight='bold')
        ax.set_xlabel(xl, color="#8bb6d6", fontsize=9); ax.set_ylabel(yl, color="#8bb6d6", fontsize=9)
        ax.tick_params(colors="#8bb6d6", labelsize=8)
        for spine in ax.spines.values(): spine.set_color(self.col_border)
        ax.grid(True, color=self.col_border, linestyle="--", alpha=0.5)

    def _show_stage(self, key):
        if key == "blank":
            self.lbl_desc.configure(text="Ready to start.")
            self.viz_player.clear("")
            return

        desc = {
            "generation": "Stars are randomly generated in the search space.",
            "running": "Black Hole algorithm is absorbing stars.",
            "finish": "Optimization finished."
        }
        self.lbl_desc.configure(text=desc.get(key, ""))
        
        paths = self.stage_images.get(key, "")
        if isinstance(paths, str):
            paths = [os.path.join(self.images_dir, paths)]
        else:
            paths = [os.path.join(self.images_dir, p) for p in paths]
            
        self.viz_player.show(paths, animate=True)

    def _get_sys_data(self):
        s = self.system_var.get()
        if s == "Ball & Beam": return init_ballandbeam({"m": safe_float(self.entries["bb_m"].get()), "R": safe_float(self.entries["bb_R"].get()), "g": safe_float(self.entries["bb_g"].get()), "L": safe_float(self.entries["bb_L"].get()), "d": safe_float(self.entries["bb_d"].get()), "J": safe_float(self.entries["bb_J"].get()), "step_amp": safe_float(self.entries["bb_s"].get())}), cost_ballandbeam
        if s == "Cruise Control": return init_cruise({"m": safe_float(self.entries["cr_m"].get()), "b": safe_float(self.entries["cr_b"].get()), "step_amp": safe_float(self.entries["cr_s"].get())}), cost_cruise
        if s == "Motor Speed Control": return init_motor({"J": safe_float(self.entries["mo_J"].get()), "b": safe_float(self.entries["mo_b"].get()), "K": safe_float(self.entries["mo_K"].get()), "R": safe_float(self.entries["mo_R"].get()), "L": safe_float(self.entries["mo_L"].get()), "step_amp": safe_float(self.entries["mo_s"].get())}), cost_motor
        if s == "Custom TF": return init_custom_tf({"num": parse_coefficients(self.entries["ctf_n"].get()), "den": parse_coefficients(self.entries["ctf_d"].get()), "step_amp": safe_float(self.entries["ctf_s"].get())}), cost_custom

    def run_opt(self):
        if self.is_running: return
        sys_n, c_type = self.system_var.get(), self.controller_var.get()
        if sys_n == "Ball & Beam" and c_type == "PI": return messagebox.showerror("Error", "PI controller is not suitable for Ball and Beam.")
        try:
            plant, cost_raw = self._get_sys_data()
            min_k, max_k = np.array([safe_float(self.min_kp.get()), safe_float(self.min_ki.get()), safe_float(self.min_kd.get())]), np.array([safe_float(self.max_kp.get()), safe_float(self.max_ki.get()), safe_float(self.max_kd.get())])
            if c_type == "PI": min_k = min_k[:2]; max_k = max_k[:2]
            elif c_type == "PD": min_k = np.array([min_k[0], min_k[2]]); max_k = np.array([max_k[0], max_k[2]])
            s = {"p": plant, "c_raw": cost_raw, "c_t": c_type, "n": int(safe_float(self.npop_var.get())), "m": int(safe_float(self.maxit_var.get())), "t": safe_float(self.sim_time_var.get()), "mn": min_k, "mx": max_k, "o": 1 if self.best_mode_var.get() else 2, "w_os": safe_float(self.w_os.get()), "w_rt": safe_float(self.w_rt.get()), "w_ess": safe_float(self.w_ess.get())}
        except Exception as e: return messagebox.showerror("Error", str(e))

        self.is_running, self.stop_requested, self.last_iter_plotted = True, False, 0
        self.history_rows_list = []
        self.gui_queue = queue.Queue()
        self.btn_run.configure(state="disabled"); self.btn_stop.configure(state="normal"); self.btn_save.configure(state="disabled")
        self.lbl_status.configure(text="Running", text_color="#ffd400"); self.lbl_prog.configure(text="Searching for the best solution...")
        
        self._style_ax(self.ax_cost, "Best Cost vs Iteration", "Iteration", "Cost"); self.canvas_cost.draw()
        self._style_ax(self.ax_sys, "System Response (Best Solution)", "Time (s)", "Amplitude"); self.canvas_sys.draw()
        for i in self.tree.get_children(): self.tree.delete(i)
        for m in self.metrics.values(): m.configure(text="--")

        self._show_stage("generation")
        self.after(1500, lambda: self._show_stage("running") if self.is_running else None)
        
        self._process_queue_job = self.after(50, self._process_queue)
        threading.Thread(target=self._worker, args=(s,), daemon=True).start()

    def stop_opt(self): self.stop_requested = True; self.lbl_status.configure(text="Stopping...", text_color="#ff4d5d")
    def reset_results(self):
        if not self.is_running:
            self.last_history = pd.DataFrame(); self.history_rows_list = []
            self.btn_save.configure(state="disabled")
            self.lbl_status.configure(text="Ready", text_color="#59ff45"); self.lbl_prog.configure(text="No run yet")
            for i in self.tree.get_children(): self.tree.delete(i)
            for m in self.metrics.values(): m.configure(text="--")
            self._style_ax(self.ax_cost, "Cost", "Iter", "Cost"); self.canvas_cost.draw()
            self._style_ax(self.ax_sys, "Response", "Time", "Amp"); self.canvas_sys.draw()
            self._show_stage("blank")

    def _worker(self, s):
        try:
            def cost_fcn(k_opt): return s["c_raw"](expand_controller_gains(k_opt, s["c_t"]), s["p"], s["c_t"], s["t"], s["o"], s["w_os"], s["w_rt"], s["w_ess"])
            def prog_fcn(h): self.gui_queue.put(h.copy())
            best, hist = black_hole_algorithm(cost_fcn, s["n"], s["m"], s["mn"], s["mx"], s["c_t"], s["p"], s["t"], lambda: self.stop_requested, prog_fcn)
            
            if self.stop_requested: self.after(0, self._finish_stop)
            else:
                t, y = simulate_feedback_response(s["p"], expand_controller_gains(best, s["c_t"]), s["c_t"], s["t"])
                self.after(0, lambda: self._finish_success(hist, t, y, s["p"].step_amp))
        except Exception as e: self.after(0, lambda: self._finish_err(str(e)))

    def _process_queue(self):
        if not self.is_running and self.gui_queue.empty(): return
        batch = []
        while not self.gui_queue.empty():
            batch.append(self.gui_queue.get_nowait())
            
        if batch:
            self.history_rows_list.extend(batch)
            for r in batch:
                curr = int(r["Iteration"])
                self.tree.insert("", "end", values=(curr, f"{r['Kp']:.4f}", f"{r['Ki']:.4f}", f"{r['Kd']:.4f}", f"{r['Cost']:.6g}", f"{r['Overshoot']:.2f}", f"{r['Rise Time']:.3f}", f"{r['Steady-State Error']:.3f}"))
            
            self.tree.yview_moveto(1)
            r = batch[-1]
            curr = int(r["Iteration"])
            
            self.metrics["Kp"].configure(text=f"{r['Kp']:.4f}"); self.metrics["Ki"].configure(text=f"{r['Ki']:.4f}"); self.metrics["Kd"].configure(text=f"{r['Kd']:.4f}")
            self.metrics["Cost"].configure(text=f"{r['Cost']:.6g}"); self.metrics["Overshoot"].configure(text=f"{r['Overshoot']:.2f}")
            self.metrics["Rise Time"].configure(text=f"{r['Rise Time']:.3f}"); self.metrics["ESS"].configure(text=f"{r['Steady-State Error']:.3f}")
            
            self.lbl_prog.configure(text=f"Iteration {curr} completed")
            
            # Update cost plot on every iteration
            df = pd.DataFrame(self.history_rows_list)
            self.ax_cost.clear(); self._style_ax(self.ax_cost, "Best Cost vs Iteration", "Iteration", "Cost")
            self.ax_cost.plot(df["Iteration"], df["Cost"], '-o', color="#38bdf8", markersize=3)
            if len(df) > 1: self.ax_cost.set_xlim(1, max(df["Iteration"]))
            self.fig_cost.tight_layout(pad=1.2); self.canvas_cost.draw_idle()

        if self.is_running:
            self._process_queue_job = self.after(50, self._process_queue)

    def _finish_success(self, hist, t, y, ref):
        self.is_running = False
        self._process_queue()
        self.last_history = pd.DataFrame(self.history_rows_list)
        self.btn_run.configure(state="normal"); self.btn_stop.configure(state="disabled"); self.btn_save.configure(state="normal")
        self.lbl_status.configure(text="Success", text_color="#59ff45"); self.lbl_prog.configure(text="Optimization completed.")
        self._show_stage("finish")
        self.ax_sys.clear(); self._style_ax(self.ax_sys, "System Response (Best Solution)", "Time (s)", "Amplitude")
        self.ax_sys.plot(t, y, color="#59ff45", linewidth=2, label="Response")
        self.ax_sys.axhline(ref, color="#ff4d5d", linestyle="--", label="Reference")
        self.ax_sys.legend(facecolor=self.col_panel, edgecolor=self.col_border, labelcolor="white")
        self.fig_sys.tight_layout(pad=1.2); self.canvas_sys.draw_idle()

    def _finish_stop(self):
        self.is_running = False
        self._process_queue()
        self.last_history = pd.DataFrame(self.history_rows_list)
        self.btn_run.configure(state="normal"); self.btn_stop.configure(state="disabled"); self.btn_save.configure(state="normal" if not self.last_history.empty else "disabled")
        self.lbl_status.configure(text="Stopped", text_color="#ff4d5d"); self.lbl_prog.configure(text="Process aborted."); self.viz_player.clear("Stopped")

    def _finish_err(self, err):
        self.is_running = False
        self.btn_run.configure(state="normal"); self.btn_stop.configure(state="disabled")
        self.lbl_status.configure(text="Error", text_color="#ff4d5d"); messagebox.showerror("Error", err)

    def save_results(self):
        if self.last_history.empty: return
        fdr = filedialog.askdirectory()
        if fdr:
            base = f"BHA_{self.system_var.get().replace(' ', '')}_{int(time.time())}"
            excel_path = os.path.join(fdr, f"{base}_Data.xlsx")
            cost_img = os.path.join(fdr, f"{base}_Cost.png")
            sys_img = os.path.join(fdr, f"{base}_Response.png")
            
            self.last_history.to_excel(excel_path, index=False)
            self.fig_cost.savefig(cost_img)
            self.fig_sys.savefig(sys_img)
            
            try:
                import openpyxl
                from openpyxl.drawing.image import Image as xlImage
                from openpyxl.styles import PatternFill, Font
                wb = openpyxl.load_workbook(excel_path)
                ws = wb.active
                
                img1 = xlImage(cost_img)
                img2 = xlImage(sys_img)
                
                # Data uses columns A-H, placing images at J2 and J20
                ws.add_image(img1, "J2")
                ws.add_image(img2, "J20")
                
                # Highlight the last iteration (best solution) row
                last_row = ws.max_row
                highlight_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
                bold_font = Font(bold=True)
                for col in range(1, 9):  # Columns A through H
                    cell = ws.cell(row=last_row, column=col)
                    cell.fill = highlight_fill
                    cell.font = bold_font
                
                wb.save(excel_path)
            except Exception as e:
                print(f"Failed to embed images in Excel: {e}")
                
            messagebox.showinfo("Saved", "Results and graphs successfully.")

if __name__ == "__main__":
    app = BlackHoleOptimizerApp()
    app.mainloop()